from __future__ import annotations

import os
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from excel_sync import (
    CURRENT_SHEET_TITLE,
    ExcelSyncError,
    MonthlyStockWorkbook,
    WorkbookCount,
    matches_bytes,
    month_title,
    read_counts_bytes,
    render_bytes,
)


class ExcelSyncBytesTests(unittest.TestCase):
    @staticmethod
    def months(counted: float | None = 9) -> list[dict]:
        return [{
            "month": "2026-09",
            "rows": [{
                "product_id": 7,
                "product": "CASINHA AZUL BEBÊ",
                "system_stock": 12,
                "counted": counted,
                "post_count_delta": 3 if counted is not None else 0,
                "final_stock": 12,
            }],
            "is_current": True,
        }]

    def test_render_bytes_builds_the_existing_layout_and_formulas(self):
        content = render_bytes(self.months())

        self.assertTrue(content.startswith(b"PK"))
        workbook = load_workbook(BytesIO(content), data_only=False)
        try:
            self.assertEqual(
                workbook.sheetnames[:2],
                [CURRENT_SHEET_TITLE, month_title("2026-09")],
            )
            current = workbook[CURRENT_SHEET_TITLE]
            self.assertEqual(current["A2"].value, "CASINHA AZUL BEBÊ")
            self.assertEqual(current["B2"].value, 12)
            self.assertEqual(current["B3"].value, "=SUM(B2:B2)")
            self.assertTrue(current.protection.sheet)

            month = workbook[month_title("2026-09")]
            self.assertEqual(month["C2"].value, 9)
            self.assertEqual(month["D2"].value, '=IF(C2="","",C2-B2)')
            self.assertEqual(month["E2"].value, '=IF(C2="",B2,C2+G2)')
            self.assertEqual(month["E3"].value, "=SUM(E2:E2)")
            self.assertFalse(month["C2"].protection.locked)
            self.assertTrue(month["B2"].protection.locked)
            self.assertTrue(month.column_dimensions["F"].hidden)
        finally:
            workbook.close()

    def test_read_counts_bytes_returns_the_same_count_objects_as_path_api(self):
        content = render_bytes(self.months())

        self.assertEqual(
            read_counts_bytes(content),
            [WorkbookCount(month="2026-09", product_id=7, quantity=9)],
        )

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "ESTOQUE SICRONIZADO.xlsx"
            path.write_bytes(content)
            self.assertEqual(MonthlyStockWorkbook(path).read_counts(), read_counts_bytes(content))

    def test_matches_bytes_detects_stale_stock(self):
        months = self.months()
        content = render_bytes(months)
        self.assertTrue(matches_bytes(content, months))

        workbook = load_workbook(BytesIO(content), data_only=False)
        output = BytesIO()
        try:
            workbook[CURRENT_SHEET_TITLE]["B2"] = 99
            workbook.save(output)
        finally:
            workbook.close()

        self.assertFalse(matches_bytes(output.getvalue(), months))

    def test_render_bytes_preserves_unmanaged_source_sheets(self):
        source = Workbook()
        notes = source.active
        notes.title = "ANOTAÇÕES"
        notes["A1"] = "Não remover"
        source_content = BytesIO()
        source.save(source_content)
        source.close()

        content = render_bytes(self.months(), source_content.getvalue())
        workbook = load_workbook(BytesIO(content), data_only=False)
        try:
            self.assertIn("ANOTAÇÕES", workbook.sheetnames)
            self.assertEqual(workbook["ANOTAÇÕES"]["A1"].value, "Não remover")
            self.assertEqual(workbook[month_title("2026-09")]["C2"].value, 9)
        finally:
            workbook.close()

    def test_path_write_uses_memory_renderer_and_keeps_atomic_replace(self):
        months = self.months()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "ESTOQUE SICRONIZADO.xlsx"
            path.write_bytes(render_bytes(self.months(counted=None)))
            workbook = MonthlyStockWorkbook(path)

            original_replace = os.replace
            with patch("excel_sync.os.replace", wraps=original_replace) as replace:
                result = workbook.write(months)

            replace.assert_called_once()
            self.assertEqual(
                result["sheets"],
                [CURRENT_SHEET_TITLE, month_title("2026-09")],
            )
            self.assertTrue(matches_bytes(path.read_bytes(), months))
            self.assertEqual(workbook.read_counts(), read_counts_bytes(path.read_bytes()))

    def test_invalid_xlsx_bytes_raise_the_existing_sync_error(self):
        for operation in (
            lambda: read_counts_bytes(b"not an xlsx"),
            lambda: matches_bytes(b"not an xlsx", self.months()),
            lambda: render_bytes(self.months(), b"not an xlsx"),
        ):
            with self.subTest(operation=operation):
                with self.assertRaisesRegex(
                    ExcelSyncError,
                    "Não foi possível abrir a planilha sincronizada",
                ):
                    operation()


if __name__ == "__main__":
    unittest.main()
